import csv
import json
import io
import xml.etree.ElementTree as ET
from typing import Any
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.colors import HexColor, white
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.core.config import settings

BLUE = HexColor("#003087")
APP_NAME = settings.APP_NAME

def export_pdf(title: str, columns: list[str], rows: list[list], output: io.BytesIO):
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), topMargin=15*mm, bottomMargin=15*mm)
    styles = getSampleStyleSheet()
    elements = []
    elements.append(Paragraph(f"{APP_NAME} - {title}", ParagraphStyle('T', parent=styles['Title'], fontSize=14, textColor=BLUE)))
    elements.append(Spacer(1, 5*mm))

    header_style = ParagraphStyle('H', parent=styles['Normal'], fontSize=8, textColor=white)
    cell_style = ParagraphStyle('C', parent=styles['Normal'], fontSize=7)

    data = [[Paragraph(c, header_style) for c in columns]]
    for row in rows:
        data.append([Paragraph(str(v) if v is not None else '', cell_style) for v in row])

    col_width = (landscape(A4)[0] - 30*mm) / len(columns)
    table = Table(data, colWidths=[col_width]*len(columns))
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), BLUE),
        ('TEXTCOLOR', (0,0), (-1,0), white),
        ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
        ('FONTSIZE', (0,1), (-1,-1), 7),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [white, HexColor("#F0F4F8")]),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('PADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(table)
    doc.build(elements)


def export_excel(title: str, columns: list[str], rows: list[list], output: io.BytesIO):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_fill = PatternFill(start_color="003087", end_color="003087", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    alt_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    for col_idx in range(1, len(columns)+1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 18

    wb.save(output)


def export_word(title: str, columns: list[str], rows: list[list], output: io.BytesIO):
    doc = Document()
    doc.add_heading(f'{APP_NAME} - {title}', level=1)

    table = doc.add_table(rows=1 + len(rows), cols=len(columns))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, col in enumerate(columns):
        cell = table.rows[0].cells[i]
        cell.text = col
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9)
        from docx.oxml.ns import qn
        shading = cell._element.get_or_add_tcPr()
        shading_elm = shading.makeelement(qn('w:shd'), {'w:fill': '003087', 'w:val': 'clear'})
        shading.append(shading_elm)

    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row):
            table.rows[row_idx + 1].cells[col_idx].text = str(value) if value is not None else ''

    doc.save(output)


def export_csv_data(columns: list[str], rows: list[list], output: io.StringIO):
    writer = csv.writer(output)
    writer.writerow(columns)
    for row in rows:
        writer.writerow(row)


def export_xml(title: str, columns: list[str], rows: list[list], output: io.BytesIO):
    root = ET.Element("data")
    root.set("title", title)
    for row in rows:
        item = ET.SubElement(root, "item")
        for col, val in zip(columns, row):
            field = ET.SubElement(item, col.replace(" ", "_").replace("°", "").replace("#", "N"))
            field.text = str(val) if val is not None else ""
    tree = ET.ElementTree(root)
    tree.write(output, encoding="unicode", xml_declaration=True)


def export_rtf(title: str, columns: list[str], rows: list[list]) -> str:
    rtf = r"{\rtf1\ansi\deff0"
    rtf += r"{\fonttbl{\f0 Arial;}}"
    rtf += r"\f0\fs20 "
    rtf += r"\b " + f"{APP_NAME} - {title}" + r"\b0\par\par "

    # Header
    rtf += r"\trowd"
    for i in range(len(columns)):
        rtf += f"\\cellx{(i+1)*2000}"
    rtf += r"\intbl "
    for col in columns:
        rtf += r"\b " + col + r"\b0\cell "
    rtf += r"\row "

    # Rows
    for row in rows:
        rtf += r"\trowd"
        for i in range(len(columns)):
            rtf += f"\\cellx{(i+1)*2000}"
        rtf += r"\intbl "
        for val in row:
            rtf += str(val if val is not None else '') + r"\cell "
        rtf += r"\row "

    rtf += "}"
    return rtf


def export_json_data(columns: list[str], rows: list[list]) -> list[dict]:
    return [dict(zip(columns, row)) for row in rows]
