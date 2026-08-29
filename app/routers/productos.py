import io
from typing import Any
from decimal import Decimal

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, or_
from sqlalchemy.orm import selectinload
from app.core.database import get_db
from app.models.producto import Producto
from app.models.proveedor import Proveedor
from app.models.laboratorios import Laboratorio
from app.models.user import User
from app.models.movimiento_stock import MovimientoStock
from app.models.deposito import Deposito
from app.models.stock_producto_deposito import StockProductoDeposito
from app.services import stock_service
from app.schemas.producto import (
    ProductoCreate, ProductoUpdate, ProductoResponse, ProductoPublicResponse,
    PreciosBulkUpdate, PreciosPorcentajeUpdate
)
from app.schemas.movimiento_stock import StockOperacionRequest
from app.services.pricing_service import (
    LISTAS,
    MARGEN_FIELDS,
    PRECIO_FIELDS,
    aplicar_a_producto,
    calcular_costos,
    calcular_precio,
    calcular_todo,
    computar_para_update,
    precios_de,
)
from app.utils.deps import get_current_user, require_role
from datetime import datetime, timezone
from app.utils.filters import apply_column_filters


def _stocks_por_deposito(p: Producto) -> list[dict]:
    """Stock del producto por depósito, ordenado como el maestro de depósitos.

    `total_blisters` es lo vendible expresado en la unidad chica: es contra eso
    que el formulario de venta decide si alcanza, sin tener que rehacer la cuenta
    de cajas × blisters_por_caja en el front.
    """
    por_caja = p.get_blisters_por_caja
    rows = [
        {
            "deposito_id": s.deposito_id,
            "nombre": s.deposito.nombre if s.deposito else None,
            "orden": s.deposito.orden if s.deposito else 0,
            "activo": s.deposito.activo if s.deposito else True,
            "cajas": s.cajas,
            "blisters": s.blisters,
            "reservado_cajas": s.reservado_cajas,
            "reservado_blisters": s.reservado_blisters,
            "total_blisters": s.total_blisters(por_caja),
            "reservado_total_blisters": s.total_reservado_blisters(por_caja),
        }
        for s in (getattr(p, "stocks_deposito", None) or [])
    ]
    rows.sort(key=lambda r: (r["orden"], r["deposito_id"]))
    return rows


# Total de cajas del producto sumando todos los depósitos. Se usa para ordenar y
# filtrar la grilla, donde interesa "cuánto hay" y no en qué góndola está.
_STOCK_TOTAL_CAJAS = (
    select(func.coalesce(func.sum(StockProductoDeposito.cajas), 0))
    .where(StockProductoDeposito.producto_id == Producto.id)
    .correlate(Producto)
    .scalar_subquery()
)


def _producto_to_response(p: Producto) -> dict:
    """Build response dict with relationship names and computed fields."""
    d = ProductoResponse.model_validate(p).model_dump()
    d["proveedor_nombre"] = p.proveedor.nombre if p.proveedor else None
    d["laboratorio_nombre"] = p.laboratorio.nombre if p.laboratorio else None
    d["stocks"] = _stocks_por_deposito(p)
    return d


router = APIRouter()


@router.get("")
async def list_productos(
    search: str = Query("", description="Buscar por nombre o código"),
    column_filters: str | None = Query(None, alias="filters", description="JSON column filters"),
    proveedor_id: int | None = Query(None),
    laboratorio_id: int | None = Query(None),
    sort_by: str = Query("nombre", description="Campo de orden: nombre, codigo, created_at, pvp"),
    sort_dir: str = Query("asc", pattern="^(asc|desc)$"),
    solo_nuevos: bool = Query(False, description="Solo productos cargados recientemente"),
    nuevos_dias: int = Query(30, ge=1, le=365, description="Ventana de días para 'producto nuevo'"),
    sin_pvp: bool = Query(False, description="Solo productos sin PVP cargado"),
    all: bool = Query(False, description="Devolver todos los productos sin paginar (para grilla en memoria)"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    filters = []

    if search:
        like_pattern = f"%{search}%"
        filters.append(
            or_(
                Producto.nombre.ilike(like_pattern),
                Producto.codigo.ilike(like_pattern),
            )
        )

    if proveedor_id:
        filters.append(Producto.proveedor_id == proveedor_id)

    if laboratorio_id:
        filters.append(Producto.laboratorio_id == laboratorio_id)

    # Herramienta comercial: productos nuevos (cargados en los últimos N días).
    if solo_nuevos:
        from datetime import datetime, timedelta
        filters.append(Producto.created_at >= datetime.now() - timedelta(days=nuevos_dias))

    # Trabajo de carga de precios: productos sin PVP.
    if sin_pvp:
        filters.append(Producto.pvp.is_(None))

    where_clause = and_(*filters) if filters else True

    base_query = (
        select(Producto)
        .outerjoin(Laboratorio, Producto.laboratorio_id == Laboratorio.id)
        .where(where_clause)
    )

    base_query = apply_column_filters(
        base_query,
        Producto,
        column_filters,
        allowed_columns={
            "codigo", "nombre", "stock_total_cajas", "stock_minimo_cajas", "status",
            "categoria_producto", "presentacion", "pvp", "costo_mas_iibb",
            "laboratorio_nombre",
            *PRECIO_FIELDS,
            *MARGEN_FIELDS,
        },
        extra_mappings={
            "laboratorio_nombre": Laboratorio.nombre,
            "stock_total_cajas": _STOCK_TOTAL_CAJAS,
        },
    )

    count_query = select(func.count()).select_from(base_query.subquery())
    total_result = await db.execute(count_query)
    total = total_result.scalar_one()

    # Orden configurable. "codigo" y "created_at" permiten ver el último producto
    # cargado (created_at desc) o recorrer los códigos en orden.
    _SORT_COLUMNS = {
        "id": Producto.id,
        "nombre": Producto.nombre,
        "codigo": Producto.codigo,
        "created_at": Producto.created_at,
        "pvp": Producto.pvp,
        "presentacion": Producto.presentacion,
        "categoria_producto": Producto.categoria_producto,
        "status": Producto.status,
        "stock_total_cajas": _STOCK_TOTAL_CAJAS,
    }
    sort_col = _SORT_COLUMNS.get(sort_by, Producto.nombre)
    order_expr = sort_col.desc() if sort_dir == "desc" else sort_col.asc()

    query = (
        base_query
        .options(
            selectinload(Producto.proveedor),
            selectinload(Producto.laboratorio)
        )
        .order_by(order_expr)
    )
    if not all:
        query = query.offset((page - 1) * page_size).limit(page_size)
    result = await db.execute(query)
    productos = result.scalars().unique().all()

    return {
        "items": [_producto_to_response(p) for p in productos],
        "total": total,
        "page": 1 if all else page,
        "page_size": total if all else page_size,
    }


async def build_public_catalogo(
    db: AsyncSession,
    *,
    search: str = "",
    lista: str = "minorista",
    laboratorio_id: int | None = None,
    page: int = 1,
    page_size: int = 25,
    all: bool = False,
) -> dict:
    """Arma el catálogo público (productos activos + precios del grupo `lista`).

    Reutilizable por el endpoint abierto `/productos/public` y por el catálogo
    con token en `/publico/lista/{token}`. Sólo expone los precios del grupo
    pedido: quien abre el catálogo minorista no recibe mayoristas ni comercio.
    """
    filters = [
        Producto.activo == True,
        Producto.status == "activo",
    ]

    if search:
        like_pattern = f"%{search}%"
        filters.append(
            or_(
                Producto.nombre.ilike(like_pattern),
                Producto.codigo.ilike(like_pattern),
                Producto.categoria_producto.ilike(like_pattern),
            )
        )

    if laboratorio_id:
        filters.append(Producto.laboratorio_id == laboratorio_id)

    where_clause = and_(*filters)

    # Count total
    count_query = select(func.count(Producto.id)).where(where_clause)
    total_result = await db.execute(count_query)
    total = total_result.scalar_one()

    # Orden: laboratorio (nulls al final) y luego nombre
    lab_order = func.coalesce(Laboratorio.nombre, "￿")
    query = (
        select(Producto)
        .outerjoin(Laboratorio, Producto.laboratorio_id == Laboratorio.id)
        .options(selectinload(Producto.laboratorio))
        .where(where_clause)
        .order_by(func.lower(lab_order), func.lower(Producto.nombre))
    )
    if not all:
        query = query.offset((page - 1) * page_size).limit(page_size)

    result = await db.execute(query)
    productos = result.scalars().all()

    def _to_item(p: Producto) -> dict:
        d = ProductoPublicResponse.model_validate(p).model_dump()
        d["laboratorio_nombre"] = p.laboratorio.nombre if p.laboratorio else None
        # En el catálogo solo importa si hay o no: se suman los depósitos activos.
        d["stock_total_cajas"] = sum(
            st.cajas for st in (p.stocks_deposito or []) if not st.deposito or st.deposito.activo
        )
        d["precios"] = precios_de(p, lista)
        return d

    return {
        "items": [_to_item(p) for p in productos],
        "total": total,
        "page": 1 if all else page,
        "page_size": total if all else page_size,
    }


@router.get("/public")
async def list_productos_public(
    search: str = Query("", description="Buscar por nombre o código"),
    lista: str = Query(
        "minorista",
        pattern="^(minorista|mayorista|comercio)$",
        description="Grupo de lista de precios a exponer",
    ),
    laboratorio_id: int | None = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100),
    all: bool = Query(False, description="Devolver todos los productos sin paginar"),
    db: AsyncSession = Depends(get_db),
):
    """Public version of product list for clients (abierto, sin token)."""
    return await build_public_catalogo(
        db, search=search, lista=lista, laboratorio_id=laboratorio_id,
        page=page, page_size=page_size, all=all,
    )


@router.get("/import-template")
async def descargar_import_template(
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    """Download a blank XLSX template for bulk product import.

    Las columnas de stock se arman con los depósitos que existen ahora mismo, que
    es exactamente lo que lee el importador. Antes la plantilla traía las columnas
    viejas ("Stock A / Stock B") y el importador ya no las reconocía: quien la
    bajaba y cargaba stock no veía ningún error, simplemente no se importaba nada.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos"

    depositos_all = (
        await db.execute(
            select(Deposito).where(Deposito.activo.is_(True)).order_by(Deposito.orden, Deposito.id)
        )
    ).scalars().all()
    stock_cols = [
        f"Stock {dep.nombre} {unidad}" for dep in depositos_all for unidad in ("cajas", "blisters")
    ]

    required_cols = ["Código", "Nombre", "Presentación", "URL PVP (Alfabeta)", "Categoría Producto"]
    optional_cols = [
        "Laboratorio", "Proveedor", "PVP", "Blisters por caja",
        *stock_cols,
        "Margen Minorista %", "Margen Mayorista %", "Margen Comercio %",
    ]
    headers = required_cols + optional_cols

    blue_fill = PatternFill(start_color="003087", end_color="003087", fill_type="solid")
    grey_fill = PatternFill(start_color="B0B0B0", end_color="B0B0B0", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = blue_fill if col_idx <= len(required_cols) else grey_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    # Fila de ejemplo. Los márgenes de comercio van vacíos: sin margen, el producto
    # no se publica en esa lista (es opt-in por producto).
    example = [
        "ABC-001", "Ibuprofeno 400mg", "Caja x 20 comp.", "https://alfabeta.net/...", "OTC",
        "Laboratorio Ejemplo", "Proveedor Ejemplo", 1500.0, 10,
        *([10, 0] * len(depositos_all)),
        30.0, 20.0, None,
    ]
    for col_idx, val in enumerate(example, 1):
        cell = ws.cell(row=2, column=col_idx, value=val)
        cell.border = thin_border

    anchos = {
        "Código": 14, "Nombre": 40, "Presentación": 20, "URL PVP (Alfabeta)": 35,
        "Categoría Producto": 20, "Laboratorio": 22, "Proveedor": 22, "PVP": 12,
        "Blisters por caja": 18,
    }
    for col_idx, name in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = anchos.get(name, 22)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_productos.xlsx"},
    )


@router.get("/precios-excel")
async def descargar_precios_excel(
    filtro: str = Query("todos"),
    filtro_id: int | None = Query(None),
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    """Download an XLSX with current prices for editing."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    query = (
        select(Producto)
        .options(selectinload(Producto.proveedor))
        .order_by(Producto.nombre)
    )
    if filtro == "proveedor" and filtro_id:
        query = query.where(Producto.proveedor_id == filtro_id)

    result = await db.execute(query)
    productos = result.scalars().unique().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Precios"

    # Las columnas de precio se generan desde LISTAS, y el import las lee por nombre
    # de encabezado. Agregar una lista nueva no rompe los archivos ya descargados.
    headers = ["ID", "Código", "Nombre", "Proveedor", "PVP Actual", "Nuevo PVP"]
    for lista in LISTAS:
        headers.append(f"{lista.excel_label} Actual")
        headers.append(f"Nuevo {lista.excel_label}")

    header_fill = PatternFill(start_color="003087", end_color="003087", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    def _editable(row_idx: int, col_idx: int):
        cell = ws.cell(row=row_idx, column=col_idx, value=None)
        cell.border = thin_border
        cell.fill = yellow_fill

    for row_idx, p in enumerate(productos, 2):
        ws.cell(row=row_idx, column=1, value=p.id).border = thin_border
        ws.cell(row=row_idx, column=2, value=p.codigo).border = thin_border
        ws.cell(row=row_idx, column=3, value=p.nombre).border = thin_border
        ws.cell(row=row_idx, column=4, value=p.proveedor.nombre if p.proveedor else "").border = thin_border
        ws.cell(row=row_idx, column=5, value=float(p.pvp) if p.pvp else None).border = thin_border
        _editable(row_idx, 6)

        col_idx = 7
        for lista in LISTAS:
            actual = getattr(p, lista.precio_field, None)
            ws.cell(row=row_idx, column=col_idx, value=float(actual) if actual else None).border = thin_border
            _editable(row_idx, col_idx + 1)
            col_idx += 2

    widths = [8, 14, 40, 18, 18, 18] + [24] * (2 * len(LISTAS))
    for col_idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=precios_productos.xlsx"},
    )


@router.get("/{id}")
async def get_producto(
    id: int,
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(get_current_user),
):
    result = await db.execute(
        select(Producto).where(Producto.id == id)
        .options(selectinload(Producto.proveedor))
    )
    producto = result.scalar_one_or_none()
    if producto is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Producto no encontrado",
        )
    return _producto_to_response(producto)


@router.post("", response_model=ProductoResponse, status_code=status.HTTP_201_CREATED)
async def create_producto(
    body: ProductoCreate,
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    existing = await db.execute(
        select(Producto).where(Producto.codigo == body.codigo)
    )
    if existing.scalar_one_or_none() is not None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Ya existe un producto con ese código",
        )

    data = body.model_dump()
    margenes = {campo: data.get(campo) for campo in MARGEN_FIELDS}
    data.update(calcular_todo(data.get("pvp"), data.get("costo_porcentaje"), margenes))
    producto = Producto(**data)
    db.add(producto)
    await db.flush()
    result = await db.execute(
        select(Producto).where(Producto.id == producto.id)
        .options(selectinload(Producto.proveedor))
    )
    producto = result.scalar_one()
    await db.commit()
    return _producto_to_response(producto)


@router.put("/{id}")
async def update_producto(
    id: int,
    body: ProductoUpdate,
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    result = await db.execute(select(Producto).where(Producto.id == id))
    producto = result.scalar_one_or_none()
    if producto is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Producto no encontrado",
        )

    update_data = body.model_dump(exclude_unset=True)

    # Los precios de venta son derivados: se recalculan desde pvp + costo% + márgenes.
    # Si el body toca alguno de esos, lo que venga en precio_venta_* se pisa.
    recalculado = computar_para_update(producto, update_data)
    if recalculado:
        for lista in LISTAS:
            update_data.pop(lista.precio_field, None)
        update_data.update(recalculado)
        # Una lista sin margen no tiene precio: hay que borrarlo si el margen se borró.
        for lista in LISTAS:
            if lista.precio_field not in recalculado:
                update_data[lista.precio_field] = None

    if "codigo" in update_data and update_data["codigo"] != producto.codigo:
        existing = await db.execute(
            select(Producto).where(Producto.codigo == update_data["codigo"])
        )
        if existing.scalar_one_or_none() is not None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Ya existe un producto con ese código",
            )

    for field, value in update_data.items():
        setattr(producto, field, value)

    await db.flush()
    result2 = await db.execute(
        select(Producto).where(Producto.id == id)
        .options(selectinload(Producto.proveedor))
    )
    producto = result2.scalar_one()
    await db.commit()
    return _producto_to_response(producto)


@router.post("/actualizar-precios-bulk")
async def actualizar_precios_bulk(
    body: PreciosBulkUpdate,
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    """Bulk update prices for multiple products."""
    updated = 0
    for item in body.items:
        result = await db.execute(select(Producto).where(Producto.id == item.producto_id))
        producto = result.scalar_one_or_none()
        if producto:
            if item.pvp is not None:
                producto.pvp = item.pvp
            for lista in LISTAS:
                valor = getattr(item, lista.precio_field, None)
                if valor is not None:
                    setattr(producto, lista.precio_field, valor)
            if item.costo_mas_iibb is not None:
                producto.costo_mas_iibb = item.costo_mas_iibb
            updated += 1

    await db.commit()
    return {"message": f"{updated} productos actualizados", "updated": updated}


@router.post("/actualizar-precios-porcentaje")
async def actualizar_precios_porcentaje(
    body: PreciosPorcentajeUpdate,
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    """
    Apply a percentage update to pvp (multiplicative) or margins/costs (additive).

    Un producto con el margen destino en NULL se saltea: no tiene esa lista y no hay
    que inventársela. Sin esto, un "+5% al margen comercio blisteado / todos" le
    crearía precio de comercio a todo el catálogo.
    """
    f_id = str(body.filtro_id) if body.filtro_id is not None else None
    query = select(Producto)
    if body.filtro == "proveedor" and f_id:
        query = query.where(Producto.proveedor_id == int(f_id))
        
    elif body.filtro == "categoria" and f_id:
        if f_id == "1":
            query = query.where(Producto.categoria_producto.ilike("OTC"))
        elif f_id == "2":
            query = query.where(Producto.categoria_producto.ilike("GENERICO"))
        else:
            return {"message": "Categoría no reconocida", "updated": 0}
    elif body.filtro == "laboratorio" and f_id:
        query = query.where(Producto.laboratorio_id == int(f_id))

    result = await db.execute(query)
    productos = result.scalars().all()

    if not productos:
        return {"message": "No se encontraron productos para los criterios", "updated": 0}

    porc_decimal = Decimal(str(body.porcentaje))
    factor_mult = Decimal(1) + (porc_decimal / Decimal(100))

    actualizados = 0
    for p in productos:
        # A. Aplicar el cambio sobre el campo destino
        if body.campo == "pvp":
            if not p.pvp:
                continue
            p.pvp = (p.pvp * factor_mult).quantize(Decimal("0.01"))
            p.fecha_act_pvp = datetime.now(timezone.utc)
        else:
            actual = getattr(p, body.campo)
            if actual is None and body.campo in MARGEN_FIELDS:
                # El producto no tiene esta lista configurada: no se le crea.
                continue
            setattr(p, body.campo, (actual or Decimal(0)) + porc_decimal)

        # B. Recalcular costos y precios derivados
        aplicar_a_producto(p)
        actualizados += 1

    await db.commit()
    return {"message": f"{actualizados} productos actualizados", "updated": actualizados}


@router.post("/precios-excel")
async def importar_precios_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    """Upload an XLSX with updated prices.

    Las columnas se leen POR NOMBRE de encabezado (no por posición), así que el
    archivo puede tener columnas nuevas o reordenadas sin romper la importación.
    Sólo se requiere la columna "id"; el resto de los "Nuevo ..." son opcionales.
    """
    import openpyxl

    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="El archivo debe ser XLSX")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active

    header_row = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: idx for idx, name in enumerate(header_row)}

    if "id" not in col:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail='El archivo no tiene la columna "ID". Descargue la plantilla desde "Exportar precios".',
        )

    def _cell(row_vals: tuple, header: str):
        idx = col.get(header)
        if idx is None or idx >= len(row_vals):
            return None
        v = row_vals[idx]
        return v if v != "" else None

    # "Nuevo <lista>" -> campo del modelo. Los headers son los mismos que genera el
    # export (excel_label), así que un archivo bajado antes de este cambio sigue
    # importando bien.
    NUEVOS = {f"nuevo {lista.excel_label.lower()}": lista.precio_field for lista in LISTAS}
    NUEVOS["nuevo pvp"] = "pvp"

    updated = 0
    skipped = 0
    errors: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or not _cell(row, "id"):
            skipped += 1
            continue

        producto_id = _cell(row, "id")
        nuevos = {campo: _cell(row, header) for header, campo in NUEVOS.items()}

        if all(v in (None, "", 0) for v in nuevos.values()):
            skipped += 1
            continue

        try:
            result = await db.execute(select(Producto).where(Producto.id == int(producto_id)))
            producto = result.scalar_one_or_none()
            if producto:
                cargados = {
                    campo: Decimal(str(valor)).quantize(Decimal("0.01"))
                    for campo, valor in nuevos.items()
                    if valor not in (None, "", 0)
                }
                for campo, valor in cargados.items():
                    setattr(producto, campo, valor)

                # Si cambió el PVP hay que rehacer los derivados: costo_neto, costo+IIBB
                # y los precios de cada lista. Los precios que la fila trae escritos a
                # mano son un override deliberado y se respetan; el resto se recalcula
                # desde su margen (y una lista sin margen sigue sin precio).
                if "pvp" in cargados:
                    costos = calcular_costos(producto.pvp, producto.costo_porcentaje)
                    if costos:
                        producto.costo_neto = costos["costo_neto"]
                        producto.costo_mas_iibb = costos["costo_mas_iibb"]
                        for lista in LISTAS:
                            if lista.precio_field in cargados:
                                continue
                            setattr(
                                producto,
                                lista.precio_field,
                                calcular_precio(
                                    costos["costo_mas_iibb"],
                                    getattr(producto, lista.margen_field),
                                ),
                            )
                updated += 1
            else:
                errors.append(f"Fila {row_idx}: Producto ID {producto_id} no encontrado")
                skipped += 1
        except Exception as e:
            errors.append(f"Fila {row_idx}: {str(e)}")
            skipped += 1

    await db.commit()
    return {"updated": updated, "skipped": skipped, "errors": errors[:20]}


@router.post("/{id}/operacion-stock")
async def operacion_stock(
    id: int,
    req: StockOperacionRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Ajuste manual o transferencia de stock entre depósitos.

    Ya no existe FRACTION: el stock se guarda en blísters y se re-normaliza a
    cajas + sueltos en cada movimiento, así que romper una caja no es una
    operación, pasa solo cuando una venta lo necesita.
    """
    result = await db.execute(select(Producto).where(Producto.id == id).with_for_update())
    producto = result.scalar_one_or_none()
    if not producto:
        raise HTTPException(status_code=404, detail="Producto no encontrado")

    tipo = (req.tipo_operacion or "").upper()
    if tipo == "MANUAL":  # nombre viejo del ajuste manual
        tipo = "ADJUST"

    if tipo == "ADJUST":
        dep = await stock_service.validar_deposito(db, req.deposito_id)
        await stock_service.ajustar(db, producto, dep.id, req.cantidad_cajas, req.cantidad_blisters)
        # Un ajuste positivo entra al depósito; uno negativo sale de él.
        entra = (req.cantidad_cajas + req.cantidad_blisters) >= 0
        origen_id, destino_id = (None, dep.id) if entra else (dep.id, None)
    elif tipo == "TRANSFER":
        origen = await stock_service.validar_deposito(db, req.deposito_id)
        destino = await stock_service.validar_deposito(db, req.deposito_destino_id)
        await stock_service.transferir(
            db, producto, origen.id, destino.id, req.cantidad_cajas, req.cantidad_blisters
        )
        origen_id, destino_id = origen.id, destino.id
    else:
        raise HTTPException(
            status_code=400, detail=f"Tipo de operación '{req.tipo_operacion}' no reconocida"
        )

    db.add(
        MovimientoStock(
            producto_id=producto.id,
            usuario_id=current_user.id,
            tipo_operacion=tipo,
            deposito_origen_id=origen_id,
            deposito_destino_id=destino_id,
            cantidad_cajas=req.cantidad_cajas,
            cantidad_blisters=req.cantidad_blisters,
            observacion=req.observacion,
        )
    )

    await db.commit()
    await db.refresh(producto)

    return _producto_to_response(producto)


@router.delete("/{id}")
async def delete_producto(
    id: int,
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    result = await db.execute(select(Producto).where(Producto.id == id))
    producto = result.scalar_one_or_none()
    if producto is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Producto no encontrado",
        )

    producto.status = "inactivo"
    producto.activo = False
    await db.commit()
    return {"message": f"Producto {producto.nombre} desactivado"}


def _map_producto_row(row: dict[str, Any]) -> dict[str, Any] | None:
    """Mapea una fila del importador JSON a campos del maestro de productos.

    El stock quedó fuera a propósito: vive por depósito y entra por ingreso de
    mercadería o ajuste manual, que dejan movimiento y responsable. Una
    importación de maestro no debería poder pisar existencias en silencio.
    """
    if "p_codigo" in row:
        pvp_str = row.get("p_precioventa", "0")
        return {
            "codigo": (row.get("p_codigo") or "").strip(),
            "nombre": (row.get("p_nombre") or "").strip(),
            "pvp": float(pvp_str) if pvp_str else None,
            "foto_url": (row.get("p_foto") or "").strip() or None,
        }
    elif "codigo" in row:
        return {
            "codigo": (row.get("codigo") or "").strip(),
            "nombre": (row.get("nombre") or "").strip(),
            "pvp": row.get("pvp") or row.get("precio_venta"),
            "foto_url": row.get("foto_url"),
        }
    return None


@router.post("/import")
async def import_productos(
    items: list[dict[str, Any]] = Body(...),
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    created = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    for i, row in enumerate(items):
        mapped = _map_producto_row(row)
        if not mapped or not mapped.get("codigo") or not mapped.get("nombre"):
            skipped += 1
            continue

        try:
            result = await db.execute(
                select(Producto).where(Producto.codigo == mapped["codigo"])
            )
            existing = result.scalar_one_or_none()

            if existing:
                for field, value in mapped.items():
                    setattr(existing, field, value)
                updated += 1
            else:
                producto = Producto(**mapped)
                db.add(producto)
                created += 1
        except Exception as e:
            errors.append(f"Fila {i + 1} ({mapped.get('codigo', '?')}): {str(e)}")
            skipped += 1

    await db.commit()

    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "total": len(items),
        "errors": errors[:20],
    }


@router.post("/import-excel")
async def import_productos_excel(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _current_user: User = Depends(require_role(["admin", "super_admin"])),
):
    """Import products from an XLSX file using the official template."""
    import openpyxl

    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="El archivo debe ser XLSX")

    contents = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active

    # Map header names (lowercase, stripped) to column index
    header_row = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: idx for idx, name in enumerate(header_row)}

    REQUIRED = {
        "código": "codigo",
        "nombre": "nombre",
        "presentación": "presentacion",
        "url pvp (alfabeta)": "url_pvp",
        "categoría producto": "categoria_producto",
    }
    missing = [h for h in REQUIRED if h not in col]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Columnas requeridas faltantes: {', '.join(missing)}",
        )

    OPTIONAL = {
        "pvp": "pvp",
        "margen minorista %": "margen_minorista",
        "margen mayorista %": "margen_mayorista",
        "margen comercio %": "margen_comercio",
    }

    # Columnas de stock: una por depósito existente, resueltas por nombre.
    # Así crear un depósito nuevo habilita su columna sin tocar código.
    # Ej. "Stock Sanalle cajas" / "Stock Sanalle blisters".
    depositos_all = (
        await db.execute(select(Deposito).order_by(Deposito.orden, Deposito.id))
    ).scalars().all()
    STOCK_COLS: dict[str, tuple[int, str]] = {}
    for dep in depositos_all:
        base = dep.nombre.strip().lower()
        STOCK_COLS[f"stock {base} cajas"] = (dep.id, "cajas")
        STOCK_COLS[f"stock {base} blisters"] = (dep.id, "blisters")

    def _cell(row_vals: tuple, header: str):
        idx = col.get(header)
        if idx is None or idx >= len(row_vals):
            return None
        v = row_vals[idx]
        return v if v != "" else None

    # Pre-load lookup caches to avoid N+1 queries
    lab_result = await db.execute(select(Laboratorio.id, Laboratorio.nombre))
    lab_cache: dict[str, int] = {n.strip().lower(): i for i, n in lab_result.all()}

    prov_result = await db.execute(select(Proveedor.id, Proveedor.nombre))
    prov_cache: dict[str, int] = {n.strip().lower(): i for i, n in prov_result.all()}

    created = 0
    updated = 0
    skipped = 0
    errors: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row or all(v is None or v == "" for v in row):
            continue

        codigo = str(_cell(row, "código") or "").strip()
        nombre = str(_cell(row, "nombre") or "").strip()
        presentacion = str(_cell(row, "presentación") or "").strip()
        url_pvp = str(_cell(row, "url pvp (alfabeta)") or "").strip()
        categoria = str(_cell(row, "categoría producto") or "").strip()

        if not all([codigo, nombre, presentacion, url_pvp, categoria]):
            errors.append(f"Fila {row_idx}: faltan campos requeridos (código, nombre, presentación, URL PVP, categoría)")
            skipped += 1
            continue

        mapped: dict[str, Any] = {
            "codigo": codigo,
            "nombre": nombre,
            "presentacion": presentacion,
            "url_pvp": url_pvp,
            "categoria_producto": categoria,
        }

        # Resolve laboratorio by name
        lab_name = str(_cell(row, "laboratorio") or "").strip()
        if lab_name:
            lab_id = lab_cache.get(lab_name.lower())
            if lab_id:
                mapped["laboratorio_id"] = lab_id
            else:
                errors.append(f"Fila {row_idx}: laboratorio '{lab_name}' no encontrado — se omitirá el campo")

        # Resolve proveedor by name
        prov_name = str(_cell(row, "proveedor") or "").strip()
        if prov_name:
            prov_id = prov_cache.get(prov_name.lower())
            if prov_id:
                mapped["proveedor_id"] = prov_id
            else:
                errors.append(f"Fila {row_idx}: proveedor '{prov_name}' no encontrado — se omitirá el campo")

        for header, field in OPTIONAL.items():
            val = _cell(row, header)
            if val is not None:
                try:
                    mapped[field] = Decimal(str(val)).quantize(Decimal("0.01"))
                except (ValueError, TypeError):
                    errors.append(f"Fila {row_idx}: valor inválido para '{header}' ({val!r})")

        # Blísters por caja: es lo que habilita vender fraccionado. Sin este dato
        # el producto entra como "solo caja" y el selector de unidad ni aparece,
        # que era justo lo que pasaba con todo lo que entraba por planilla.
        bl_por_caja = _cell(row, "blisters por caja")
        if bl_por_caja is not None:
            try:
                n = int(bl_por_caja)
                if n < 1:
                    raise ValueError
                mapped["blisters_por_caja"] = n
                # Mismo criterio que el alta manual: si fracciona, se habilita la
                # venta por blíster. Con 1 blíster por caja el "blíster" sería la
                # caja entera, así que ahí no se toca.
                if n > 1:
                    mapped["vende_blister"] = True
            except (ValueError, TypeError):
                errors.append(
                    f"Fila {row_idx}: 'Blisters por caja' tiene que ser un entero >= 1 ({bl_por_caja!r})"
                )

        # Stock por depósito de esta fila: {deposito_id: {"cajas": n, "blisters": n}}
        stock_fila: dict[int, dict[str, int]] = {}
        for header, (dep_id, campo) in STOCK_COLS.items():
            val = _cell(row, header)
            if val is None:
                continue
            try:
                stock_fila.setdefault(dep_id, {})[campo] = int(val)
            except (ValueError, TypeError):
                errors.append(f"Fila {row_idx}: valor inválido para '{header}' ({val!r})")

        try:
            result = await db.execute(select(Producto).where(Producto.codigo == codigo))
            existing = result.scalar_one_or_none()
            if existing:
                for field, value in mapped.items():
                    setattr(existing, field, value)
                # Sin esto, importar márgenes dejaba los precios de venta sin recalcular.
                aplicar_a_producto(existing)
                producto_fila = existing
                updated += 1
            else:
                nuevo = Producto(**mapped)
                aplicar_a_producto(nuevo)
                db.add(nuevo)
                # El producto nuevo necesita id antes de poder colgarle stock.
                await db.flush()
                producto_fila = nuevo
                created += 1

            for dep_id, cantidades in stock_fila.items():
                await stock_service.establecer(
                    db,
                    producto_fila,
                    dep_id,
                    cantidades.get("cajas", 0),
                    cantidades.get("blisters", 0),
                )
        except Exception as e:
            errors.append(f"Fila {row_idx} ({codigo}): {str(e)}")
            skipped += 1

    await db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "total": created + updated + skipped, "errors": errors[:20]}
